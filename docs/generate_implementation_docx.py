"""One-off script: generates docs/IROS_Implementation_and_Features.docx —
a Word document summarizing the current, actual implementation and features
of the Institutional Research Operating System (IROS), for stakeholders who
want a document rather than reading the markdown docs / source directly.

Not part of the running application. Run with a Python environment that has
`python-docx` installed (not an app dependency - see README for how this was
generated):

    python3 -m venv /tmp/docx_venv && /tmp/docx_venv/bin/pip install python-docx
    /tmp/docx_venv/bin/python docs/generate_implementation_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

OUTPUT_PATH = Path(__file__).parent / "IROS_Implementation_and_Features.docx"

ACCENT = RGBColor(0x2D, 0x5B, 0xAA)
DARK = RGBColor(0x1A, 0x1A, 0x1A)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = ACCENT
    return h


def add_bullets(doc, items, style="List Bullet"):
    for item in items:
        p = doc.add_paragraph(style=style)
        if isinstance(item, tuple):
            label, rest = item
            r = p.add_run(label)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def build() -> None:
    doc = Document()

    # Title page
    title = doc.add_heading("Institutional Research Operating System (IROS)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("Implementation & Features Reference")
    sub_run.font.size = Pt(18)
    sub_run.font.color.rgb = ACCENT
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("An 8-desk, 25-agent AI equity-research pipeline\n").italic = True
    meta.add_run("Local-first · citation-required · human-approval-gated publishing").italic = True
    doc.add_page_break()

    # 1. Executive summary
    add_heading(doc, "1. Executive Summary", 1)
    doc.add_paragraph(
        "IROS automates the end-to-end workflow of a senior institutional equity-research "
        "analyst: pulling real public filings and market data, extracting typed facts, running "
        "five specialist analysts, building and valuing a real 3-statement DCF model, and "
        "producing a single, evidence-linked investment decision that only publishes after an "
        "adversarial internal review and explicit human compliance sign-off."
    )
    add_bullets(doc, [
        ("25 specialist agents ", "organized into 8 \u201cdesks\u201d, each with its own "
         "collaboration pattern (sequential handoff, parallel wave, gated parallel, or "
         "adversarial chain)."),
        ("Two independent surfaces: ", "a standalone real-time Ticker Dashboard (any real "
         "ticker, live Yahoo Finance/SEC data, AI narrative, DCF/Excel export) and the full "
         "25-agent Research Cockpit pipeline."),
        ("Zero-hallucination discipline: ", "every claim in the final thesis must cite a "
         "resolvable evidence reference; agents that lack sufficient real evidence abstain "
         "rather than invent an answer."),
        ("Local-first: ", "runs entirely on a laptop with SQLite + a local Ollama LLM - no "
         "cloud API keys, no paid data feeds, no telemetry."),
        ("Fail-closed governance: ", "no export, memo, or trade recommendation is ever "
         "produced until a Compliance & MNPI Guardrail agent clears the run AND a human "
         "explicitly approves publication."),
    ])

    # 2. Architecture
    add_heading(doc, "2. Architecture", 1)
    doc.add_paragraph(
        "Orchestration is handled entirely by LangGraph's StateGraph - there is no separate "
        "workflow engine (e.g. Temporal). Every agent is a plain Python class extending a "
        "shared BaseAgent, registered into a single graph with declared depends_on edges that "
        "LangGraph topologically sorts into parallel \u201cwaves\u201d automatically."
    )
    add_table(
        doc,
        ["Layer", "Technology"],
        [
            ["Orchestration", "LangGraph (StateGraph), no Temporal"],
            ["Backend", "Python 3.12, FastAPI, Pydantic v2, SQLAlchemy (async)"],
            ["Database", "SQLite by default (zero-config); PostgreSQL via docker-compose.yml"],
            ["LLM routing", "ModelRouter \u2192 Ollama (llama3.1, nomic-embed-text), deterministic mock fallback"],
            ["Retrieval", "Hybrid BM25 + embedding cosine-similarity RAG over real SEC filing text"],
            ["Frontend", "Next.js 14 (App Router), React 18, TypeScript, Zustand, Recharts"],
            ["Real data sources", "yfinance (live quotes/fundamentals), sec-edgar-downloader (real filings)"],
            ["Modeling", "openpyxl (live-formula Excel export), formulas (independent re-verification)"],
            ["Testing", "pytest + Hypothesis (backend), Vitest + Playwright (frontend)"],
        ],
    )
    doc.add_paragraph(
        "Only 4 of the 25 agents call an LLM at all: Industry & Competitive Context, "
        "Transcript Intelligence, Filing Narrative Analyst, and the standing Continuous "
        "Monitoring Agent. The other 21 are pure, deterministic Python with no hallucination "
        "surface by construction."
    )

    # 3. The 8 desks
    add_heading(doc, "3. The 8-Desk / 25-Agent Pipeline", 1)
    desks = [
        ("1. Mandate & Coverage", "Sequential handoff", "Research Planner \u2192 Entity & Universe "
         "Resolver. Scopes the mandate and resolves the exact company/security + peer set before "
         "anything else can run."),
        ("2. Data Acquisition", "Parallel wave", "Acquisition Orchestrator (SEC EDGAR + news "
         "connectors) and Expert Notes Ingestion (stricter MNPI screen for uploaded material) run "
         "concurrently."),
        ("3. Document & Fact Extraction", "Gated parallel", "Document Intelligence parses raw "
         "filings first; Financial Fact Extraction and Transcript Intelligence then run against "
         "that parsed output."),
        ("4. Fundamental & Context Research", "Parallel wave", "Five specialist analysts - "
         "Financial Integrity, Industry & Competitive Context, Filing Narrative Analyst, Macro/ESG "
         "Signal Context, and Market Data Validation - all read the same frozen fact base."),
        ("5. Strategy & Idea Generation", "Parallel wave", "Risk & Catalyst (taxonomized risk map "
         "+ dated catalyst calendar) and Variant Perception Ideation (non-consensus hypotheses with "
         "mandatory falsification tests)."),
        ("6. Modeling & Valuation", "Sequential handoff", "Operating Model (driver-based 3-statement "
         "forecast) \u2192 Valuation Scenario (DCF / comps / probability-weighted scenarios)."),
        ("7. Investment Committee", "Adversarial chain (deliberately never parallelized)",
         "Thesis Synthesizer proposes \u2192 Devil's Advocate and Verification Agent independently "
         "attack/fact-check \u2192 Adjudicator weighs both and sets final confidence/publishability "
         "\u2192 Compliance & MNPI Guardrail has the final legal word."),
        ("8. Decision, Distribution & Monitoring", "Gated parallel + standing agent",
         "Portfolio Sizing Advisor and Report & Notebook Agent run concurrently once compliance "
         "clears; IC Memo Writer follows. Continuous Monitoring Agent is a standing member that "
         "runs on its own schedule, independent of any single research run, watching for the next "
         "material trigger."),
    ]
    for name, mode, desc in desks:
        p = doc.add_paragraph()
        r = p.add_run(f"{name} ")
        r.bold = True
        r2 = p.add_run(f"\u2014 {mode}")
        r2.italic = True
        r2.font.color.rgb = ACCENT
        doc.add_paragraph(desc)

    doc.add_paragraph(
        "Desk 7 is the anti-hallucination core of the whole system: it is the one place in the "
        "pipeline that is permanently, deliberately non-parallelizable, mirroring a real "
        "investment committee's separation of powers. The Adjudicator's disagreement-adjudication "
        "protocol discounts the bull thesis's own confidence for every unresolved Devil's Advocate "
        "claim and every failed verification check - disagreements are never averaged away or "
        "silently ignored."
    )

    # 4. Feature tour
    add_heading(doc, "4. Feature Tour", 1)
    features = [
        ("Home", "Command bar to start a new research run (staged review or full auto-run mode), "
         "a live backend-status indicator, and a list of recent runs across the workspace."),
        ("Ticker Dashboard", "A standalone, always-fast feature independent of the 25-agent "
         "pipeline: real-time Yahoo Finance quote/ratios/peer comps, an on-demand AI analyst "
         "summary (with a fact-check counter showing how many provided real figures were actually "
         "cited), a 207-point deterministic data-quality validation pass, multi-year financial "
         "statements, and a full DCF/3-statement financial model with a live-formula Excel export."),
        ("Research Cockpit", "The full 8-desk, 25-agent pipeline for a given run: per-desk, "
         "per-agent live status/confidence/latency, and every cockpit module (Financial Scorecard, "
         "Risk Map, Catalyst Calendar, Estimate Bridge, Bull/Bear Debate, Variant Perception, "
         "Decision Card, What Changed, Provenance & Quality, Contradictions Panel)."),
        ("Notebook", "Ask a natural-language follow-up question about any completed run; answers "
         "are grounded in and cite that run's own frozen evidence (bull thesis pillars, bear-case "
         "claims, risks, catalysts) or honestly say the question isn't answerable from the "
         "available evidence."),
        ("Portfolio", "A simple watchlist with real, live Yahoo Finance quotes for every ticker "
         "added, persisted locally in the browser."),
        ("Admin", "A live connector-health board that actually probes Yahoo Finance, SEC EDGAR, "
         "and the local Ollama instance on demand (not a simulated status), plus a manual trigger "
         "for the Continuous Monitoring poll across the workspace's watchlist."),
        ("Activity / Events", "A cross-workspace audit trail of every research run ever started, "
         "newest first, linking through to each run's cockpit."),
    ]
    for name, desc in features:
        p = doc.add_paragraph()
        r = p.add_run(f"{name}: ")
        r.bold = True
        p.add_run(desc)

    # 5. Quality & testing
    add_heading(doc, "5. Quality, Testing & Governance", 1)
    add_bullets(doc, [
        ("178 backend tests passing ", "(pytest + Hypothesis property tests) across unit, "
         "contract, e2e, calibration, adversarial, and golden-file suites, plus 2 intentional "
         "xfail cases."),
        ("Frontend test suite: ", "Vitest unit tests plus a real Playwright end-to-end test that "
         "drives an actual browser through submitting a ticker, running the full pipeline, and "
         "verifying the final decision card renders."),
        ("Deterministic, in-agent validators: ", "several agents enforce their own contracts at "
         "runtime (e.g. every_claim_has_citation, blocked_if_not_compliance_cleared, "
         "disagreement_protocol_applied), not just external test assertions."),
        ("Independent verification: ", "the Verification Agent recomputes every material number "
         "from raw facts before the Adjudicator ever sees the thesis - it never trusts a cached "
         "result."),
        ("Fail-closed compliance gate: ", "export, memo generation, and any publication-facing "
         "action are hard-blocked until compliance_report.cleared is True and a human has "
         "explicitly approved - verified via both automated tests and live manual QA."),
    ])
    doc.add_paragraph(
        "A comprehensive, live end-to-end QA pass (covering every page, every one of the 25 "
        "agents, and real runs against multiple real tickers including AAPL and Texas "
        "Instruments/TXN) was completed as part of preparing this implementation for handoff. "
        "Five real defects were found and fixed during that pass, each verified against a fresh "
        "live run afterward:"
    )
    add_table(
        doc,
        ["#", "Issue", "Fix"],
        [
            ["1", "A duplicated bear-case argument could appear twice in the Bull/Bear "
             "Debate and Adjudicator panels.", "Devil's Advocate now attacks only the first "
             "matching thesis pillar for a given argument."],
            ["2", "The built-in demo ticker could be silently routed through a real (always-"
             "failing) SEC lookup when real data sources were enabled.", "The demo ticker now "
             "always uses its fixture data, matching the convention already used by its sibling "
             "agents."],
            ["3", "The Notebook's Q&A only searched the bull thesis, so most realistic questions "
             "(risks, catalysts, bear case) returned \u201cnot answerable.\u201d", "Expanded the "
             "searchable evidence to include bear-case claims, risks, and catalysts."],
            ["4", "In fully-autonomous run mode, a viewer could be silently stranded on the "
             "first step of the pipeline even after the whole run finished.", "Replaced an "
             "unreliable WebSocket-only \u201cis it done\u201d signal with an authoritative, "
             "REST-verified one, plus a polling fallback."],
            ["5", "A data-completeness score was permanently reported as 0% for every real "
             "ticker, regardless of how much real data was actually retrieved.", "Corrected the "
             "expected-document-type list to match real SEC filing types instead of demo-only "
             "vocabulary."],
        ],
        widths=[0.4, 3.0, 3.0],
    )

    # 6. Setup & deployment
    add_heading(doc, "6. Setup & Deployment", 1)
    doc.add_paragraph("Local development (no containers required):")
    add_bullets(doc, [
        "Pull two local Ollama models once: llama3.1 and nomic-embed-text.",
        "Backend: create a Python 3.12 virtual environment in agent/, "
        "pip install -e \".[dev]\", copy .env.example to .env, run python main.py "
        "(serves on port 8000).",
        "Frontend: npm install in frontend/, copy .env.example to .env.local, "
        "npm run dev (serves on port 3000).",
    ], style="List Number")
    doc.add_paragraph(
        "Optional production-scale infrastructure (PostgreSQL, Redis, MinIO) is available via "
        "the repository's docker-compose.yml, but nothing in the default setup requires it - "
        "SQLite and local disk are the fully-supported default path."
    )
    doc.add_paragraph(
        "Deployment note: this is a stateful, full-stack application (a Python backend with its "
        "own database, plus a local LLM dependency), not a static site. See the accompanying "
        "deployment guidance for what a free-hosting deployment would and would not cover."
    )

    # 7. Known limitations
    add_heading(doc, "7. Known Limitations (Disclosed, Not Defects)", 1)
    add_bullets(doc, [
        ("Earnings-call transcripts ", "for real tickers are not yet sourced automatically "
         "(the connector is fixture-only today); the affected agent correctly abstains rather "
         "than fabricating a transcript."),
        ("DCF results for capital-intensive companies mid-cycle ", "(e.g. a semiconductor "
         "company in a multi-year fab buildout) can look extreme under the generic driver-based "
         "model, which holds trailing capex/margin ratios flat rather than modeling a taper; the "
         "system already surfaces an explicit caveat in the decision text asking the reviewer to "
         "weigh capacity-cycle positioning."),
        ("First-time filing analysis for a new ticker ", "can take a few minutes (real filings "
         "are chunked and embedded sequentially against the local LLM); results are cached for 24 "
         "hours per ticker afterward."),
        ("The Event Workspace detail page ", "(a single event's source coverage/reconciliation/"
         "Q&A view) is an intentionally deferred stub, not a regression."),
    ])

    doc.add_page_break()
    add_heading(doc, "Appendix: Where to Read More", 1)
    add_bullets(doc, [
        "docs/AGENTS_MASTER_REFERENCE.md \u2014 exhaustive, one-agent-at-a-time reference.",
        "docs/architecture.md \u2014 system architecture and data flow.",
        "docs/data_contracts.md \u2014 the typed ResearchState every agent reads/writes.",
        "docs/workflow_state_machine.md \u2014 staged vs. auto run-mode semantics.",
        "docs/source_policy.md \u2014 citable-source and MNPI/compliance policy.",
        "docs/adr/ \u2014 Architecture Decision Records.",
        "README.md \u2014 quick start and feature tour.",
    ])

    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
